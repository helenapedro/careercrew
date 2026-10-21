"""Shared document-access logic (single source of truth).

This module is the one place where input is turned into plain text, whether it
arrived as an uploaded file or as pasted text:

- ``read_resume`` reads a PDF / DOCX / TXT file confined to the uploads folder.
- ``read_resume_text`` normalizes pasted resume text into the same clean form.
- ``read_job_description`` returns pasted JD text, or reads an uploaded
  PDF / DOCX / TXT job description from the uploads folder.

The same logic backs the CrewAI tools, the FastMCP server, and the stdio MCP
client, so the three paths can never diverge.

Tool guardrails live here, next to the code they protect: uploaded files are
confined to the uploads folder (no path traversal) and pasted URLs are rejected
with a clear message instead of being stored as junk "job description" text.
"""

from __future__ import annotations

import os
import re

from careercrew.config import get_settings

MAX_RESUME_CHARS = 100_000
MAX_JOB_DESCRIPTION_CHARS = 100_000
RESUME_EXTENSIONS = {".pdf", ".docx", ".txt"}


def _upload_dir() -> str:
    return get_settings().upload_dir


class JobDescriptionError(ValueError):
    """Raised when a job description cannot be read.

    Carries a human-readable, user-facing message so the UI can show it in the
    error box instead of crashing the run.
    """


# --------------------------------------------------------------------------- #
# Resume input: file (PDF / DOCX / TXT) or pasted text
# --------------------------------------------------------------------------- #
def _clean_text(text: str) -> str:
    """Collapse whitespace so file and pasted input normalize identically."""
    return re.sub(r"\s+", " ", text or "").strip()


def _safe_upload_path(path: str, kind: str = "file") -> str:
    """Resolve a path inside the uploads folder; reject anything else.

    Any directory component (``../``, ``a/b``, absolute paths) is rejected
    outright: only a bare file name inside the uploads folder is allowed.
    """
    candidate = (path or "").strip()
    if not candidate or candidate != os.path.basename(candidate):
        raise ValueError(f"{kind} path must be a bare file name inside the uploads folder")
    base = os.path.abspath(_upload_dir())
    full = os.path.abspath(os.path.join(base, candidate))
    if not full.startswith(base + os.sep):
        raise ValueError(f"{kind} path escapes the uploads folder")
    if not os.path.isfile(full):
        raise FileNotFoundError(f"{kind} not found: {candidate}")
    return full


def _read_pdf(full_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(full_path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(full_path: str) -> str:
    import docx

    document = docx.Document(full_path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    # Tables often hold skills/dates in resumes; keep their cell text too.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def extract_text(full_path: str) -> str:
    """Extract text from a resume file: PDF, DOCX, or TXT."""
    lower = full_path.lower()
    if lower.endswith(".pdf"):
        raw = _read_pdf(full_path)
    elif lower.endswith(".docx"):
        raw = _read_docx(full_path)
    else:
        with open(full_path, encoding="utf-8", errors="replace") as handle:
            raw = handle.read()
    return _clean_text(raw)[:MAX_RESUME_CHARS]


def read_resume(path: str) -> str:
    """Read a resume file (PDF, DOCX, or TXT) and return its text.

    The path is confined to the uploads folder; anything else is rejected.
    """
    return extract_text(_safe_upload_path(path, "resume"))


def read_resume_text(text: str) -> str:
    """Normalize pasted resume text into the same clean form as a file read."""
    cleaned = _clean_text(text)
    if not cleaned:
        raise ValueError("pasted resume text is empty")
    return cleaned[:MAX_RESUME_CHARS]


# --------------------------------------------------------------------------- #
# Job description input: pasted text or an uploaded file
# --------------------------------------------------------------------------- #
def _looks_like_url(source: str) -> bool:
    return bool(re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", source))


def read_job_description(source: str) -> str:
    """Return JD text from an uploaded file (PDF/DOCX/TXT) or pasted text.

    A bare file name is read from the uploads folder through the same
    :func:`extract_text` path used for resumes (so PDF, DOCX, and TXT are all
    supported). Anything else is treated as pasted JD text. URLs are rejected
    with a clear message: job boards block automated access, so we ask the user
    to download the posting and upload the file, or paste the text.
    """
    source = (source or "").strip()
    if not source:
        raise JobDescriptionError("No job description was provided.")
    if _looks_like_url(source):
        raise JobDescriptionError(
            "Job description URLs are not read directly (job boards block "
            "automated access). Download the posting and upload the file, or paste "
            "the job description text instead."
        )
    if os.path.basename(source) == source and _has_supported_extension(source):
        try:
            text = extract_text(_safe_upload_path(source, "job description"))
        except FileNotFoundError:
            # Not an uploaded file after all; fall through to treat it as text.
            return _clean_text(source)[:MAX_JOB_DESCRIPTION_CHARS]
        if not text:
            raise JobDescriptionError(
                "The uploaded job description file had no readable text. "
                "Paste the job description text instead."
            )
        return text[:MAX_JOB_DESCRIPTION_CHARS]
    return _clean_text(source)[:MAX_JOB_DESCRIPTION_CHARS]


def _has_supported_extension(name: str) -> bool:
    lower = (name or "").lower()
    return any(lower.endswith(ext) for ext in RESUME_EXTENSIONS)
